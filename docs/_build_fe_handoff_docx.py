"""Build FE Hostinger Google Sign-In + issues Word document."""
from docx import Document
from docx.shared import Pt, Inches
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from pathlib import Path

OUT = Path(__file__).with_name("Frontend_Hostinger_Google_SignIn_Handoff.docx")


def set_run_font(run, size=11, bold=False, code=False):
    run.bold = bold
    run.font.size = Pt(size)
    font_name = "Consolas" if code else "Calibri"
    run.font.name = font_name
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    r_fonts.set(qn("w:ascii"), font_name)
    r_fonts.set(qn("w:hAnsi"), font_name)


def add_para(doc, text, bold=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, bold=bold)
    p.paragraph_format.space_after = Pt(6)
    return p


def add_mixed(doc, parts):
    """parts: list of (text, bold)."""
    p = doc.add_paragraph()
    for text, bold in parts:
        r = p.add_run(text)
        set_run_font(r, bold=bold)
    p.paragraph_format.space_after = Pt(6)
    return p


def add_bullet(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style="List Bullet")
    if bold_prefix:
        r1 = p.add_run(bold_prefix)
        set_run_font(r1, bold=True)
        r2 = p.add_run(text)
        set_run_font(r2)
    else:
        r = p.add_run(text)
        set_run_font(r)
    return p


def add_code(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "F5F5F5")
    shd.set(qn("w:val"), "clear")
    p.paragraph_format.element.get_or_add_pPr().append(shd)
    for line in text.strip("\n").split("\n"):
        run = p.add_run(line + "\n")
        set_run_font(run, size=9, code=True)
    return p


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        r = cell.paragraphs[0].add_run(h)
        set_run_font(r, bold=True, size=10)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = table.rows[ri + 1].cells[ci]
            cell.text = ""
            r = cell.paragraphs[0].add_run(str(val))
            set_run_font(r, size=10)
    doc.add_paragraph()
    return table


def main():
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.85)
        section.right_margin = Inches(0.85)

    title = doc.add_heading(
        "NexusLexis Frontend — Issues + Required Fixes", 0
    )
    for run in title.runs:
        set_run_font(run, size=20, bold=True)

    add_para(doc, "For: Frontend team  |  From: Backend  |  Rechecked: 20 Oct 2026")
    add_para(
        doc,
        "Site: https://nexuslexis.law (Hostinger)   |   Auth API: "
        "https://nexus-lexis-backend-45v4.vercel.app/api/auth",
    )
    add_mixed(
        doc,
        [
            ("Status: ", True),
            (
                "Auth backend is working. Live Hostinger frontend is not using the required Google flow "
                "and Hostinger still 404s SPA routes. Client secrets will not fix this.",
                False,
            ),
        ],
    )

    doc.add_heading("1. Live check (20 Oct 2026) — what is broken", 1)
    add_table(
        doc,
        ["Check", "Result", "Meaning"],
        [
            [
                "Auth /ping + /health",
                "200, db ok",
                "Backend is up",
            ],
            [
                "Auth GET /google/start",
                "302 to Google",
                "Redirect flow on backend works",
            ],
            [
                "https://nexuslexis.law/",
                "200",
                "Home works",
            ],
            [
                "https://nexuslexis.law/login",
                "404",
                "No SPA rewrite on Hostinger",
            ],
            [
                "https://nexuslexis.law/signup",
                "404",
                "No SPA rewrite — signup URL itself is broken",
            ],
            [
                "https://nexuslexis.law/google-auth.html",
                "404",
                "Google return page not uploaded",
            ],
            [
                "https://nexuslexis.netlify.app/login",
                "200",
                "Netlify SPA rewrite is fine; Hostinger is not",
            ],
            [
                "Live JS bundle index-BAz4ttsT.js",
                "GoogleOAuthProvider + POST /google/token",
                "Old GIS popup still live. No /google/start in the build",
            ],
        ],
    )
    add_mixed(
        doc,
        [
            ("Bottom line: ", True),
            (
                "Changing Google Client ID/secret on FE does not fix Hostinger. "
                "The live Hostinger deploy is still the old popup + token API, and Apache is not serving React routes.",
                False,
            ),
        ],
    )

    doc.add_heading("2. Issue A — Google Sign-Up stuck on Hostinger", 1)
    doc.add_heading("Why it fails", 2)
    add_bullet(
        doc,
        "Live Hostinger app still uses GIS popup (GoogleOAuthProvider) then POST /api/auth/google/token. "
        "On Hostinger this often becomes FETCH_ERROR (“Unable to reach the server”) or Google origin_mismatch.",
    )
    add_bullet(
        doc,
        "Hostinger Apache has no SPA fallback. Direct URLs /login /signup /account return Hostinger 404 HTML.",
    )
    add_bullet(
        doc,
        "Auth redirect now returns to /google-auth.html?token=… That file is not on Hostinger, so even the correct "
        "OAuth return lands on 404.",
    )
    add_bullet(
        doc,
        "Client secret is backend-only. FE must never receive or embed GOOGLE_CLIENT_SECRET. "
        "FE only uses Client ID if GIS is used (we do not want GIS on Hostinger).",
    )

    doc.add_heading("3. What FE must do (required)", 1)
    add_para(doc, "Do all three. Secrets alone are not enough.", bold=True)

    doc.add_heading("Step 1 — Upload two files to Hostinger public_html", 2)
    add_para(
        doc,
        "Upload next to index.html (not inside /assets).",
    )
    add_table(
        doc,
        ["File to upload", "Put on Hostinger as"],
        [
            ["docs/hostinger/.htaccess", "public_html/.htaccess"],
            ["docs/hostinger/google-auth.html", "public_html/google-auth.html"],
        ],
    )
    add_para(doc, "Or download from live Auth API:")
    add_bullet(
        doc,
        "https://nexus-lexis-backend-45v4.vercel.app/api/auth/google/hostinger/.htaccess",
    )
    add_bullet(
        doc,
        "https://nexus-lexis-backend-45v4.vercel.app/api/auth/google/hostinger/google-auth.html",
    )
    add_para(doc, "What they do:")
    add_bullet(
        doc,
        " — Apache rewrite so /login /signup /account serve index.html instead of 404.",
        bold_prefix=".htaccess",
    )
    add_bullet(
        doc,
        " — After Google, Auth redirects here with tokens. Page writes sessionStorage then goes to /.",
        bold_prefix="google-auth.html",
    )
    add_para(doc, "Verify after upload (hard refresh):")
    add_code(
        doc,
        """https://nexuslexis.law/login              → 200 (app shell, NOT Hostinger 404 page)
https://nexuslexis.law/signup             → 200
https://nexuslexis.law/google-auth.html   → 200 (“Signing you in…”)""",
    )

    doc.add_heading("Step 2 — Change the Google button (redeploy this build to Hostinger)", 2)
    add_para(
        doc,
        "Remove GIS popup on Hostinger production: GoogleOAuthProvider / @react-oauth/google as the primary path. "
        "Do not POST /google/token as the main Hostinger flow.",
    )
    add_para(doc, "Use full-page redirect:")
    add_code(
        doc,
        """const AUTH =
  import.meta.env.VITE_AUTH_API_URL ||
  'https://nexus-lexis-backend-45v4.vercel.app/api/auth';

export function startGoogleSignIn(role = 'client') {
  const safeRole = ['client', 'lawyer', 'ca'].includes(role) ? role : 'client';
  fetch(`${AUTH}/ping`).catch(() => {});
  const returnTo = encodeURIComponent(window.location.origin);
  window.location.href =
    `${AUTH}/google/start?role=${encodeURIComponent(safeRole)}&returnTo=${returnTo}`;
}

// Login + SignUp buttons:
<button type="button" onClick={() => startGoogleSignIn(selectedRole || 'client')}>
  Continue with Google
</button>""",
    )
    add_para(doc, "On SignUp, pass selected role: client | lawyer | ca.")
    add_para(doc, "Optional warm-up on app boot:")
    add_code(
        doc,
        """useEffect(() => {
  fetch(`${AUTH}/ping`).catch(() => {});
}, []);""",
    )

    doc.add_heading("Step 3 — After redirect, hydrate the session", 2)
    add_para(
        doc,
        "google-auth.html already writes the same keys as current tokenStorage:",
    )
    add_table(
        doc,
        ["sessionStorage key", "Value"],
        [
            ["accessToken", "JWT"],
            ["token", "same JWT (alias)"],
            ["refreshToken", "refresh token"],
            ["user", "cleared — refill with GET /me"],
        ],
    )
    add_para(
        doc,
        "On app load, if accessToken exists, call GET {AUTH}/me with Authorization: Bearer <accessToken> "
        "and put the user into Redux / auth context.",
    )

    doc.add_heading("4. End-to-end Google flow (after FE change)", 1)
    add_code(
        doc,
        """User clicks Continue with Google
  → GET {AUTH}/google/start?role=client&returnTo=https://nexuslexis.law
  → Google consent
  → GET {AUTH}/google/callback?code=…
  → 302 https://nexuslexis.law/google-auth.html?token=…&refreshToken=…
  → page writes sessionStorage
  → redirect /
  → React sees accessToken → GET /me → logged in""",
    )
    add_para(doc, "Error return:")
    add_code(doc, "/google-auth.html?googleError=<message>&next=/login")

    doc.add_heading("5. Google Cloud Console (FE + backend shared client)", 1)
    add_para(
        doc,
        "OAuth type: Web application. FE uses Client ID only. Secret stays on Vercel Auth.",
    )
    add_para(doc, "Authorized JavaScript origins (no trailing slash):")
    for origin in [
        "https://nexuslexis.law",
        "https://www.nexuslexis.law",
        "https://nexuslexis.netlify.app",
        "http://localhost:5173",
    ]:
        add_bullet(doc, origin)
    add_para(doc, "Authorized redirect URI (exact — Auth callback, NOT the website):")
    add_bullet(
        doc,
        "https://nexus-lexis-backend-45v4.vercel.app/api/auth/google/callback",
    )
    add_para(
        doc,
        "If you rotated a new Client ID, put the same ID in FE env AND Vercel GOOGLE_CLIENT_ID. "
        "Put the new secret only in Vercel GOOGLE_CLIENT_SECRET. Do not put the secret in the frontend repo or Hostinger files.",
    )
    add_para(
        doc,
        "Live Hostinger bundle Client ID still: 925512874310-duojhjmbplliooe0mqlgm532fogmq40g.apps.googleusercontent.com",
    )

    doc.add_heading("6. Issue B — Lawyer rejected but FE shows Pending / cannot re-apply", 1)
    add_para(
        doc,
        "Backend already sets verification to Rejected and allows re-apply. FE must read these fields from "
        "login, refresh, and GET /me (not a cached JWT-only Pending flag).",
    )
    add_code(
        doc,
        """{
  "user": {
    "verificationStatus": "Rejected",
    "canAccessDashboard": false,
    "nextStep": "complete_lawyer_application",
    "pendingVerificationRole": null,
    "canReapply": true,
    "canApplyAsLawyer": true,
    "message": "Your lawyer application was rejected. You can update your details and apply for verification again."
  },
  "accountAccess": { "same fields" }
}""",
    )
    add_table(
        doc,
        ["Condition", "UI"],
        [
            [
                "nextStep === verification_pending OR pendingVerificationRole is set",
                "Pending verification screen ONLY",
            ],
            [
                "verificationStatus === Rejected OR canReapply === true",
                "Show Rejected + open lawyer apply form again",
            ],
            [
                "canApplyAsLawyer === true",
                "Enable POST /profile/lawyer/apply (re-apply)",
            ],
        ],
    )
    add_para(
        doc,
        "Do not treat Rejected as Pending. Do not hide apply just because a lawyer_profile row already exists. "
        "Same pattern for CA: canApplyAsCA / canReapply.",
    )

    doc.add_heading("7. Auth API cheat sheet", 1)
    add_para(doc, "Base: https://nexus-lexis-backend-45v4.vercel.app/api/auth")
    add_table(
        doc,
        ["Method", "Path", "Use"],
        [
            ["GET", "/ping", "Warm-up"],
            ["GET", "/config", "googleAuth.recommendedMode = redirect"],
            ["GET", "/google/start?role=&returnTo=", "PRIMARY Google entry on Hostinger"],
            ["GET", "/google/callback", "Google hits this — FE must not call it"],
            ["POST", "/google/token", "Legacy GIS — do not use as primary on Hostinger"],
            ["GET", "/me", "User + verification / re-apply flags"],
            ["POST", "/login", "Email/password"],
            ["POST", "/refresh", "Body { refreshToken }"],
            ["POST", "/profile/lawyer/apply", "Lawyer apply / re-apply after reject"],
        ],
    )

    doc.add_heading("8. Acceptance checklist", 1)
    for item in [
        ".htaccess uploaded to Hostinger public_html",
        "google-auth.html uploaded to Hostinger public_html",
        "/login and /signup return 200 on nexuslexis.law (hard refresh)",
        "google-auth.html returns 200",
        "New JS bundle on Hostinger (not still index-BAz4ttsT.js)",
        "Google button uses window.location.href to /google/start — no GIS popup on Hostinger",
        "Google signup completes and lands logged in",
        "GET /me after Google landing hydrates user",
        "Lawyer rejected in admin → FE shows Rejected (not Pending) and apply works again",
        "Google Console origins + redirect URI match section 5",
        "Client secret is NOT in frontend env or repo",
    ]:
        p = doc.add_paragraph(style="List Bullet")
        r = p.add_run("☐  " + item)
        set_run_font(r)

    doc.add_heading("9. Troubleshooting", 1)
    add_table(
        doc,
        ["Symptom", "Cause / fix"],
        [
            [
                "Error 400: origin_mismatch",
                "Add https://nexuslexis.law and https://www.nexuslexis.law as JS origins",
            ],
            [
                "Unable to reach the server / FETCH_ERROR",
                "Stop using POST /google/token on Hostinger; use /google/start",
            ],
            [
                "/login or /signup is Hostinger 404",
                "Upload .htaccess; confirm it is in public_html",
            ],
            [
                "Google finishes then 404",
                "Upload google-auth.html; Auth now redirects there, not /login",
            ],
            [
                "Stuck on Signing you in…",
                "URL must include token=; check sessionStorage keys",
            ],
            [
                "Logged in but no user name",
                "Call GET /me after bridge redirect",
            ],
            [
                "Redirect URI mismatch",
                "Only Auth callback URL in Google Console (section 5)",
            ],
            [
                "Still old popup after deploy",
                "Hostinger still serving old bundle index-BAz4ttsT.js — clear CDN/cache and re-upload dist",
            ],
        ],
    )

    doc.add_heading("10. Files from backend repo", 1)
    add_table(
        doc,
        ["Path", "Purpose"],
        [
            ["docs/hostinger/.htaccess", "SPA rewrite for Hostinger"],
            ["docs/hostinger/google-auth.html", "Token bridge page"],
            ["docs/hostinger/README.md", "Short upload note"],
            [
                "docs/Frontend_Hostinger_Google_SignIn_Handoff.docx",
                "This Word document",
            ],
        ],
    )

    doc.save(OUT)
    print(f"Wrote {OUT} ({OUT.stat().st_size} bytes)")


if __name__ == "__main__":
    main()
